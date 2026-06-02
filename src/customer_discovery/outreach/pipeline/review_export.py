from __future__ import annotations

import csv
import re
from pathlib import Path

from customer_discovery.models.final import FinalBrief
from customer_discovery.models.outreach import OutreachPack

EXCEL_COLUMNS = [
    "rank",
    "company_id",
    "company_name",
    "website",
    "final_score",
    "confidence",
    "final_stage",
    "ready_to_send",
    "contact_persona",
    "contact_name",
    "contact_title",
    "contact_email",
    "contact_linkedin",
    "contact_source",
    "hook_email",
    "hook_linkedin",
    "email_subject",
    "email_body",
    "linkedin_connection_note",
    "discovery_question",
    "company_summary",
    "pain_points",
    "value_props_for_them",
    "disqualifier_notes",
    "review_warnings",
    "evidence_ids_used",
    "important_urls",
]

REVIEW_COLUMNS = [
    "rank",
    "company_id",
    "company_name",
    "website",
    "final_score",
    "confidence",
    "ready_to_send",
    "contact_name",
    "contact_title",
    "contact_linkedin",
    "contact_source",
    "email_subject",
    "email_preview",
    "pack_file",
    "review_warnings",
]


def _one_line(text: str | None, max_len: int = 200) -> str:
    if not text:
        return ""
    flat = re.sub(r"\s+", " ", text.strip())
    if len(flat) <= max_len:
        return flat
    return flat[: max_len - 3] + "..."


def _pack_filename(p: OutreachPack) -> str:
    rank = p.rank if p.rank is not None else 0
    safe_id = re.sub(r"[^\w.-]", "_", p.company_id)
    return f"{rank:03d}_{safe_id}.md"



def _signal_claims(signals, limit: int = 5) -> list[str]:
    return [s.claim for s in signals[:limit] if s.claim]


def _render_outreach_decision(
    p: OutreachPack,
    brief: "FinalBrief | None" = None,
    lead_row: dict[str, str] | None = None,
) -> list[str]:
    lines = ["", "## Should we reach out?", ""]

    pros: list[str] = []
    cons: list[str] = []

    pros.extend(p.value_props_for_them)
    pros.extend(p.pain_points)
    if brief:
        pros.extend(_signal_claims(brief.positive_signals))
    if lead_row:
        for part in (lead_row.get("top_positive_signals") or "").split(";"):
            part = part.strip()
            if part and part not in pros:
                pros.append(part)

    cons.extend(p.disqualifier_notes)
    cons.extend(p.review_warnings)
    if brief:
        cons.extend(brief.disqualifiers)
        cons.extend(_signal_claims(brief.negative_signals))
        if brief.manual_review_required:
            cons.append("Manual review flagged during research")
    if lead_row:
        for part in (lead_row.get("top_negative_signals") or "").split(";"):
            part = part.strip()
            if part and part not in cons:
                cons.append(part)
        for part in (lead_row.get("disqualifiers") or "").split(";"):
            part = part.strip()
            if part and part not in cons:
                cons.append(part)

    # dedupe preserving order
    def dedupe(items: list[str]) -> list[str]:
        seen: set[str] = set()
        out: list[str] = []
        for x in items:
            if x and x not in seen:
                seen.add(x)
                out.append(x)
        return out

    pros = dedupe(pros)
    cons = dedupe(cons)

    lines.append("### Pros")
    if pros:
        lines.extend(f"- {x}" for x in pros[:8])
    else:
        lines.append("- _(none listed)_")

    lines.append("")
    lines.append("### Cons")
    if cons:
        lines.extend(f"- {x}" for x in cons[:8])
    else:
        lines.append("- _(none listed)_")

    return lines


def render_outreach_pack(
    p: OutreachPack,
    *,
    brief: FinalBrief | None = None,
    lead_row: dict[str, str] | None = None,
) -> str:
    """Human-readable view of one outreach_packs.jsonl entry (not the queue CSV)."""
    lines = [
        f"# {p.company_name}",
        "",
        "_From `outreach_packs.jsonl` — outreach-specific fields (hooks, value props, drafts)._",
        "",
        "## Overview",
        "",
        f"- **Company ID:** `{p.company_id}`",
        f"- **Rank:** {p.rank}",
        f"- **Score:** {p.final_score} ({p.confidence})",
        f"- **Research stage:** {p.final_stage}",
        f"- **Ready to send:** {p.ready_to_send}",
        f"- **Website:** {p.website or ''}",
    ]
    if p.generated_at:
        lines.append(f"- **Generated:** {p.generated_at.isoformat()}")
    if p.review_warnings:
        lines.append(f"- **Warnings:** {', '.join(p.review_warnings)}")

    fit_label = brief.fit_label if brief else ""
    if fit_label:
        lines.append(f"- **Fit label:** {fit_label}")

    lines.extend(_render_outreach_decision(p, brief, lead_row))

    if p.hook_email or p.hook_linkedin:
        lines.extend(["", "## Hooks", ""])
        if p.hook_email:
            lines.extend(["", "### Email hook", "", p.hook_email.strip()])
        if p.hook_linkedin:
            lines.extend(["", "### LinkedIn hook", "", p.hook_linkedin.strip()])

    lines.extend(["", "## Email draft", "", f"**Subject:** {p.email_subject or ''}", ""])
    if p.email_body:
        lines.extend(["```", p.email_body.strip(), "```", ""])

    lines.extend(["", "## LinkedIn connection note", ""])
    if p.linkedin_connection_note:
        lines.append(p.linkedin_connection_note.strip())
    else:
        lines.append("_(none)_")

    if p.discovery_question:
        lines.extend(["", "## Discovery question", "", p.discovery_question.strip()])

    if p.company_summary:
        lines.extend(["", "## Company summary", "", p.company_summary.strip()])

    if p.pain_points:
        lines.extend(["", "## Pain points", ""])
        lines.extend(f"- {x}" for x in p.pain_points)

    if p.value_props_for_them:
        lines.extend(["", "## Value props (for them)", ""])
        lines.extend(f"- {x}" for x in p.value_props_for_them)

    if p.disqualifier_notes:
        lines.extend(["", "## Disqualifier notes", ""])
        lines.extend(f"- {x}" for x in p.disqualifier_notes)

    lines.extend(["", "## Contact", ""])
    c = p.contact
    if c.name:
        lines.append(f"- **Name:** {c.name}")
    if c.title:
        lines.append(f"- **Title:** {c.title}")
    if c.persona:
        lines.append(f"- **Persona:** {c.persona}")
    if c.email:
        lines.append(f"- **Email:** {c.email}")
    if c.linkedin_url:
        lines.append(f"- **LinkedIn:** {c.linkedin_url}")
    lines.append(f"- **Source:** {c.contact_source}")

    if p.evidence_ids_used:
        lines.extend(["", "## Evidence IDs used", ""])
        lines.extend(f"- `{eid}`" for eid in p.evidence_ids_used)

    if p.important_urls:
        lines.extend(["", "## Important URLs", ""])
        for u in p.important_urls:
            mark = " *(used in reasoning)*" if u.used_in_reasoning else ""
            lines.append(f"- **[{u.source_type}]({u.url})**{mark}")
            if u.why_important:
                lines.append(f"  - {u.why_important}")

    return "\n".join(lines) + "\n"


def _bullet_lines(items: list[str]) -> str:
    return "\n".join(f"• {x}" for x in items if x)


def _format_urls(p: OutreachPack) -> str:
    lines: list[str] = []
    for u in p.important_urls:
        mark = " *" if u.used_in_reasoning else ""
        lines.append(f"[{u.source_type}] {u.url}{mark}")
        if u.why_important:
            lines.append(f"  {u.why_important}")
    return "\n".join(lines)


def _pack_excel_row(p: OutreachPack) -> dict[str, str | int | bool]:
    c = p.contact
    return {
        "rank": p.rank or "",
        "company_id": p.company_id,
        "company_name": p.company_name,
        "website": p.website or "",
        "final_score": p.final_score,
        "confidence": p.confidence,
        "final_stage": p.final_stage,
        "ready_to_send": p.ready_to_send,
        "hook_email": (p.hook_email or "").strip(),
        "hook_linkedin": (p.hook_linkedin or "").strip(),
        "contact_persona": c.persona or "",
        "contact_name": c.name or "",
        "contact_title": c.title or "",
        "contact_email": c.email or "",
        "contact_linkedin": c.linkedin_url or "",
        "contact_source": c.contact_source,
        "email_subject": p.email_subject or "",
        "email_body": (p.email_body or "").strip(),
        "linkedin_connection_note": (p.linkedin_connection_note or "").strip(),
        "discovery_question": (p.discovery_question or "").strip(),
        "company_summary": (p.company_summary or "").strip(),
        "pain_points": _bullet_lines(p.pain_points),
        "value_props_for_them": _bullet_lines(p.value_props_for_them),
        "disqualifier_notes": _bullet_lines(p.disqualifier_notes),
        "review_warnings": "; ".join(p.review_warnings),
        "evidence_ids_used": ", ".join(p.evidence_ids_used),
        "important_urls": _format_urls(p),
    }


def write_outreach_excel(path: Path, packs: list[OutreachPack]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wrap = Alignment(wrap_text=True, vertical="top")
    sorted_packs = sorted(packs, key=lambda p: p.rank or 9999)

    wb = Workbook()
    ws = wb.active
    ws.title = "Outreach"

    ws.append(EXCEL_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    long_cols = {
        "hook_email",
        "hook_linkedin",
        "email_body",
        "linkedin_connection_note",
        "company_summary",
        "pain_points",
        "value_props_for_them",
        "disqualifier_notes",
        "important_urls",
    }
    long_col_idxs = {EXCEL_COLUMNS.index(c) + 1 for c in long_cols}

    for p in sorted_packs:
        row_data = _pack_excel_row(p)
        ws.append([row_data[c] for c in EXCEL_COLUMNS])
        row_idx = ws.max_row
        for col_idx in long_col_idxs:
            ws.cell(row=row_idx, column=col_idx).alignment = wrap

    ws.freeze_panes = "A2"
    if sorted_packs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(EXCEL_COLUMNS))}{len(sorted_packs) + 1}"

    widths = {
        "rank": 6,
        "company_id": 22,
        "company_name": 18,
        "website": 28,
        "final_score": 10,
        "confidence": 12,
        "final_stage": 12,
        "ready_to_send": 14,
        "hook_email": 40,
        "hook_linkedin": 36,
        "contact_persona": 24,
        "contact_name": 18,
        "contact_title": 22,
        "contact_email": 24,
        "contact_linkedin": 36,
        "contact_source": 14,
        "email_subject": 36,
        "email_body": 56,
        "linkedin_connection_note": 40,
        "discovery_question": 40,
        "company_summary": 48,
        "pain_points": 40,
        "value_props_for_them": 40,
        "disqualifier_notes": 32,
        "review_warnings": 24,
        "evidence_ids_used": 28,
        "important_urls": 48,
    }
    for idx, col_name in enumerate(EXCEL_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col_name, 16)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def export_outreach_review(
    output_dir: Path,
    packs: list[OutreachPack],
    *,
    briefs_path: Path | None = None,
    leads_path: Path | None = None,
) -> dict[str, Path]:
    """Readable exports from outreach_packs.jsonl (markdown + Excel + summary CSV)."""
    from customer_discovery.storage.staged_jsonl import index_by_company

    briefs_map: dict[str, FinalBrief] = {}
    if briefs_path and briefs_path.exists():
        briefs_map = index_by_company(briefs_path, FinalBrief)

    lead_rows: dict[str, dict[str, str]] = {}
    if leads_path and leads_path.exists():
        with leads_path.open(encoding="utf-8") as lf:
            reader = csv.DictReader(lf)
            for row in reader:
                cid = (row.get("company_id") or "").strip()
                if cid:
                    lead_rows[cid] = row

    packs_dir = output_dir / "packs"
    packs_dir.mkdir(parents=True, exist_ok=True)
    review_path = output_dir / "outreach_review.csv"

    sorted_packs = sorted(packs, key=lambda p: p.rank or 9999)
    rows: list[dict[str, str]] = []

    for p in sorted_packs:
        pack_name = _pack_filename(p)
        pack_path = packs_dir / pack_name
        pack_path.write_text(
            render_outreach_pack(
                p,
                brief=briefs_map.get(p.company_id),
                lead_row=lead_rows.get(p.company_id),
            ),
            encoding="utf-8",
        )
        rows.append(
            {
                "rank": str(p.rank or ""),
                "company_id": p.company_id,
                "company_name": p.company_name,
                "website": p.website or "",
                "final_score": str(p.final_score),
                "confidence": p.confidence,
                "ready_to_send": str(p.ready_to_send).lower(),
                "contact_name": p.contact.name or "",
                "contact_title": p.contact.title or "",
                "contact_linkedin": p.contact.linkedin_url or "",
                "contact_source": p.contact.contact_source,
                "email_subject": p.email_subject or "",
                "email_preview": _one_line(p.email_body),
                "pack_file": f"packs/{pack_name}",
                "review_warnings": "; ".join(p.review_warnings),
            }
        )

    with review_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=REVIEW_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    excel_path = output_dir / "outreach_packs.xlsx"
    write_outreach_excel(excel_path, sorted_packs)

    readme_path = output_dir / "README.md"
    _write_outreach_index(readme_path, sorted_packs)

    return {
        "review_csv": review_path,
        "review_xlsx": excel_path,
        "packs_dir": packs_dir,
        "readme": readme_path,
    }


def _write_outreach_index(path: Path, packs: list[OutreachPack]) -> None:
    lines = [
        "# Outreach packs (readable)",
        "",
        f"{len(packs)} entries from **`outreach_packs.jsonl`** — hooks, email/LinkedIn drafts,",
        "value props, and outreach contact info. (Not the same as `outreach_queue.csv`.)",
        "",
        "| Rank | Company | Ready | Read pack |",
        "| ---: | --- | :---: | --- |",
    ]
    for p in packs:
        pack_rel = f"packs/{_pack_filename(p)}"
        ready = "yes" if p.ready_to_send else "no"
        lines.append(f"| {p.rank or ''} | {p.company_name} | {ready} | [{pack_rel}]({pack_rel}) |")
    lines.extend(
        [
            "",
            "## Files",
            "",
            "- **`packs/`** — one markdown file per `outreach_packs.jsonl` row (start here)",
            "- `outreach_packs.xlsx` — same data in Excel",
            "- `outreach_review.csv` — short index only",
            "- `outreach_queue.csv` — flat send queue (subset of columns)",
            "- `outreach_packs.jsonl` — source JSONL",
            "",
            f"Regenerate: `customer-discovery outreach export --output-dir {path.parent}`",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
